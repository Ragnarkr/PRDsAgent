from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable


@dataclass(slots=True)
class Section:
    title: str
    level: int
    content: str


@dataclass(slots=True)
class ParsedDocument:
    source_path: Path
    content: str
    sections: list[Section] = field(default_factory=list)
    metadata: dict[str, str] = field(default_factory=dict)


class DocumentParser:
    """Parse supported PRD files into a unified structure."""

    def __init__(self) -> None:
        self._parsers: dict[str, Callable[[Path], ParsedDocument]] = {
            ".md": self._parse_markdown,
            ".markdown": self._parse_markdown,
            ".txt": self._parse_text,
            ".docx": self._parse_docx,
        }

    def parse(self, file_path: str | Path) -> ParsedDocument:
        path = Path(file_path)
        suffix = path.suffix.lower()
        parser = self._parsers.get(suffix)
        if parser is None:
            supported = ", ".join(sorted(self._parsers.keys()))
            raise ValueError(f"Unsupported file type: {suffix}. Supported: {supported}")
        if not path.exists():
            raise FileNotFoundError(path)
        return parser(path)

    def _parse_markdown(self, path: Path) -> ParsedDocument:
        text = _read_text_with_fallback(path)
        sections = _split_markdown_sections(text)
        return ParsedDocument(
            source_path=path,
            content=text,
            sections=sections,
            metadata={"format": "markdown"},
        )

    def _parse_text(self, path: Path) -> ParsedDocument:
        text = _read_text_with_fallback(path)
        sections = [Section(title="全文", level=1, content=text.strip())]
        return ParsedDocument(
            source_path=path,
            content=text,
            sections=sections,
            metadata={"format": "text"},
        )

    def _parse_docx(self, path: Path) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError("python-docx is required to parse .docx files") from exc

        document = Document(str(path))
        lines: list[str] = []
        sections: list[Section] = []
        current_title = "全文"
        current_level = 1
        buffer: list[str] = []

        def flush_buffer() -> None:
            if buffer:
                sections.append(
                    Section(
                        title=current_title,
                        level=current_level,
                        content="\n".join(buffer).strip(),
                    )
                )
                buffer.clear()

        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            lines.append(text)

            style_name = (para.style.name or "").lower()
            is_heading = style_name.startswith("heading") or style_name.startswith("标题")
            if is_heading:
                flush_buffer()
                current_title = text
                current_level = _extract_heading_level(style_name)
            else:
                buffer.append(text)

        flush_buffer()
        full_text = "\n".join(lines)
        if not sections:
            sections = [Section(title="全文", level=1, content=full_text)]

        return ParsedDocument(
            source_path=path,
            content=full_text,
            sections=sections,
            metadata={"format": "docx"},
        )


def _read_text_with_fallback(path: Path) -> str:
    for encoding in ("utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _split_markdown_sections(text: str) -> list[Section]:
    sections: list[Section] = []
    current_title = "全文"
    current_level = 1
    buffer: list[str] = []

    def flush_buffer() -> None:
        if buffer:
            sections.append(
                Section(
                    title=current_title,
                    level=current_level,
                    content="\n".join(buffer).strip(),
                )
            )
            buffer.clear()

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            if title:
                flush_buffer()
                current_title = title
                current_level = level
                continue
        buffer.append(line)

    flush_buffer()
    if not sections:
        sections = [Section(title="全文", level=1, content=text.strip())]
    return sections


def _extract_heading_level(style_name: str) -> int:
    digits = "".join(char for char in style_name if char.isdigit())
    if digits:
        return int(digits)
    return 1
