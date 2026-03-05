from __future__ import annotations

import shutil
import unittest
import uuid
from pathlib import Path

from prds_agent.parsers import DocumentParser


def _has_python_docx() -> bool:
    try:
        import docx  # noqa: F401

        return True
    except ImportError:
        return False


class DocumentParserTest(unittest.TestCase):
    def setUp(self) -> None:
        self.parser = DocumentParser()
        self.root = Path("tests") / ".tmp" / f"parser_{uuid.uuid4().hex}"
        self.root.mkdir(parents=True, exist_ok=True)

    def tearDown(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    def test_parse_markdown(self) -> None:
        path = self.root / "sample.md"
        path.write_text("# 标题\n\n内容\n## 验收标准\n响应时间<=3秒\n", encoding="utf-8")

        parsed = self.parser.parse(path)
        self.assertEqual(parsed.metadata["format"], "markdown")
        self.assertGreaterEqual(len(parsed.sections), 2)

    def test_parse_text(self) -> None:
        path = self.root / "sample.txt"
        path.write_text("纯文本需求文档", encoding="utf-8")

        parsed = self.parser.parse(path)
        self.assertEqual(parsed.metadata["format"], "text")
        self.assertEqual(len(parsed.sections), 1)

    @unittest.skipUnless(_has_python_docx(), "python-docx not installed")
    def test_parse_docx(self) -> None:
        from docx import Document

        path = self.root / "sample.docx"
        doc = Document()
        doc.add_heading("需求背景", level=1)
        doc.add_paragraph("这里是背景内容")
        doc.save(path)

        parsed = self.parser.parse(path)
        self.assertEqual(parsed.metadata["format"], "docx")
        self.assertGreaterEqual(len(parsed.sections), 1)


if __name__ == "__main__":
    unittest.main()
